"""Turn uploaded files into plain text (documents) or rows (spreadsheets)."""

import csv
import hashlib
import io
import os
import re
import zipfile
from html.parser import HTMLParser
from xml.etree import ElementTree
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import PurePosixPath

DOCUMENT_EXTS = {".docx", ".pdf", ".pptx", ".txt", ".md", ".markdown", ".html", ".htm", ".rtf", ".odt"}
SPREADSHEET_EXTS = {".xlsx", ".xlsm", ".csv"}
ARCHIVE_EXTS = {".zip"}
# Google Drive for desktop shows Docs, Sheets and Slides as these: a link to the file, with none of its text.
GOOGLE_SHORTCUTS = {".gdoc", ".gsheet", ".gslides", ".gdraw", ".gform", ".gmap", ".gsite"}
MAX_TEXT_CHARS = 600_000  # ~150k tokens; anything larger is truncated with a warning
MAX_ZIP_FILES = 2_000
MAX_ZIP_BYTES = 300 * 1024 * 1024    # unpacked, for the whole archive
MAX_ZIP_ENTRY_BYTES = 20 * 1024 * 1024


class UnsupportedFile(ValueError):
    pass


@dataclass
class ParsedDocument:
    text: str
    created: datetime | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def words(self) -> int:
        return len(self.text.split())


@dataclass
class ParsedSheet:
    name: str
    rows: list[list[str]]


def kind_of(filename: str) -> str:
    ext = PurePosixPath(filename.lower()).suffix
    if ext in DOCUMENT_EXTS:
        return "document"
    if ext in SPREADSHEET_EXTS:
        return "spreadsheet"
    if ext in ARCHIVE_EXTS:
        return "archive"
    if ext in GOOGLE_SHORTCUTS:
        return "google_shortcut"
    return "unsupported"


def _skippable(inner: str) -> bool:
    name = os.path.basename(inner)
    return inner.startswith("__MACOSX/") or "/__MACOSX/" in inner or name.startswith(("._", "~$")) or name in (".DS_Store", "desktop.ini")


def expand_archives(files: list[tuple[str, bytes]]) -> list[tuple[str, bytes, str | None]]:
    """Unpacks .zip files (a Google Drive folder download, a Takeout export) into the files inside them, each keeping
    its folder path under where the zip was. Returns (path, data, why it can't be read or None)."""
    out: list[tuple[str, bytes, str | None]] = []
    for path, data in files:
        if kind_of(path) != "archive":
            out.append((path, data, None))
            continue
        where = os.path.dirname(path)
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                entries = [i for i in z.infolist() if not i.is_dir() and not _skippable(i.filename)]
                if not entries:
                    out.append((path, b"", "The zip file has nothing in it to read."))
                    continue
                # A Drive download often holds no folder of its own: the zip's name stands in for one.
                top = PurePosixPath(path).stem if all("/" not in i.filename for i in entries) else ""
                total = 0
                for n, info in enumerate(entries):
                    inner = "/".join(x for x in (where, top, info.filename) if x)
                    if n >= MAX_ZIP_FILES:
                        out.append((path, b"", f"Only the first {MAX_ZIP_FILES:,} files in the zip were read."))
                        break
                    if info.file_size > MAX_ZIP_ENTRY_BYTES:
                        out.append((inner, b"", "Larger than 20 MB."))
                        continue
                    total += info.file_size
                    if total > MAX_ZIP_BYTES:
                        out.append((path, b"", "The zip unpacks to more than 300 MB; the rest was left out. Add it in smaller parts."))
                        break
                    out.append((inner, z.read(info), None))
        except (zipfile.BadZipFile, zipfile.LargeZipFile, NotImplementedError, RuntimeError) as e:
            out.append((path, b"", f"Couldn't open the zip file ({e.__class__.__name__})."))
    return out


def text_hash(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text).strip().lower()
    return hashlib.sha256(normalized.encode()).hexdigest()


def _docx(data: bytes) -> ParsedDocument:
    import docx  # python-docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip()
                if t and (not cells or cells[-1] != t):  # merged cells repeat
                    cells.append(t)
            if cells:
                parts.append(" | ".join(cells))
    created = d.core_properties.created
    return ParsedDocument("\n".join(parts), created)


def _pdf(data: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    doc = ParsedDocument(text)
    if len(text.strip()) < 20 and len(reader.pages) > 0:
        doc.warnings.append("PDF has no extractable text (probably scanned); OCR isn't supported yet.")
    return doc


def _pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame and shape.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
        if texts:
            slides.append(f"[Slide {i}]\n" + "\n".join(texts))
    return ParsedDocument("\n\n".join(slides), prs.core_properties.created)


class _HtmlText(HTMLParser):
    """The readable text of a web page or a Google Docs "Web page" export: headings, paragraphs, list items, table
    rows, each on its own line; scripts, styles and the head left out."""
    BLOCKS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "blockquote",
              "pre", "table", "ul", "ol", "hr", "header", "footer", "dt", "dd"}
    HIDDEN = {"script", "style", "head", "title", "noscript", "template", "svg"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0
        self.title = ""
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.HIDDEN:
            self.hidden += 1
            self._in_title = tag == "title"
        elif tag in self.BLOCKS:
            self.parts.append("\n")
            if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
                self.parts.append("#" * int(tag[1]) + " ")
            elif tag == "li":
                self.parts.append("- ")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag):
        if tag in self.HIDDEN:
            self.hidden = max(0, self.hidden - 1)
            self._in_title = False
        elif tag in self.BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.hidden:
            self.parts.append(re.sub(r"\s+", " ", data))


def _html(data: bytes) -> ParsedDocument:
    raw = data.decode("utf-8", errors="replace")
    if m := re.search(rb'charset=["\']?([\w-]+)', data[:2000]):
        try:
            raw = data.decode(m.group(1).decode(), errors="replace")
        except LookupError:
            pass
    parser = _HtmlText()
    parser.feed(raw)
    text = "".join(parser.parts)
    lines = [ln.strip(" |") for ln in text.splitlines()]
    text = "\n".join(ln for ln in lines if ln.strip("-# "))
    title = parser.title.strip()
    return ParsedDocument(f"# {title}\n\n{text}" if title and title not in text[:300] else text)


_RTF_SKIP = {"fonttbl", "colortbl", "stylesheet", "info", "pict", "header", "footer", "headerl", "headerr",
             "footerl", "footerr", "listtable", "listoverridetable", "rsidtbl", "generator", "xmlnstbl", "themedata",
             "colorschememapping", "latentstyles", "datastore", "mmathPr", "object", "fldinst", "filetbl", "revtbl"}
_RTF_TOKEN = re.compile(r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\([^a-zA-Z])|([{}])|[\r\n]+|([^\\{}\r\n]+)")


def _rtf(data: bytes) -> ParsedDocument:
    """Plain text out of RTF (TextEdit, WordPad, some exports): control words become breaks or are dropped, and
    groups that hold no text (fonts, colors, pictures, …) are skipped."""
    raw = data.decode("latin-1")
    out: list[str] = []
    stack: list[tuple[bool, int]] = []
    skip, uc, pending_skip = False, 1, 0
    for m in _RTF_TOKEN.finditer(raw):
        word, arg, hexed, sym, brace, text = m.groups()
        if brace == "{":
            stack.append((skip, uc))
        elif brace == "}":
            skip, uc = stack.pop() if stack else (False, 1)
        elif sym:
            if sym == "*":
                skip = True
            elif sym in "\\{}":
                if not skip:
                    out.append(sym)
            elif sym == "~" and not skip:
                out.append(" ")
        elif word:
            if word in _RTF_SKIP:
                skip = True
            elif skip:
                continue
            elif word in ("par", "line", "sect", "page", "row"):
                out.append("\n")
            elif word == "tab" or word == "cell":
                out.append("\t" if word == "tab" else " | ")
            elif word == "uc" and arg:
                uc = int(arg)
            elif word == "u" and arg:
                code = int(arg)
                out.append(chr(code + 65536 if code < 0 else code))
                pending_skip = uc
            elif word in ("emdash", "endash"):
                out.append("—" if word == "emdash" else "–")
            elif word in ("lquote", "rquote"):
                out.append("’")
            elif word in ("ldblquote", "rdblquote"):
                out.append("”" if word == "rdblquote" else "“")
            elif word == "bullet":
                out.append("•")
        elif hexed:
            if pending_skip:
                pending_skip -= 1
            elif not skip:
                out.append(bytes([int(hexed, 16)]).decode("cp1252", errors="replace"))
        elif text and not skip:
            if pending_skip:
                drop = min(pending_skip, len(text))
                text, pending_skip = text[drop:], pending_skip - drop
            out.append(text)
    return ParsedDocument("".join(out))


_ODF = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"


def _odt_text(el) -> str:
    parts = [el.text or ""]
    for child in el:
        tag = child.tag
        if tag == f"{_ODF}s":
            parts.append(" " * int(child.get(f"{_ODF}c", "1")))
        elif tag == f"{_ODF}tab":
            parts.append("\t")
        elif tag == f"{_ODF}line-break":
            parts.append("\n")
        elif tag != f"{_ODF}note":
            parts.append(_odt_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


def _odt(data: bytes) -> ParsedDocument:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ElementTree.fromstring(z.read("content.xml"))
    lines = []
    for el in root.iter():
        if el.tag == f"{_ODF}h":
            lines.append("#" * int(el.get(f"{_ODF}outline-level", "1")) + " " + _odt_text(el))
        elif el.tag == f"{_ODF}p":
            lines.append(_odt_text(el))
    return ParsedDocument("\n".join(lines))


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    ext = PurePosixPath(filename.lower()).suffix
    if ext == ".docx":
        doc = _docx(data)
    elif ext == ".pdf":
        doc = _pdf(data)
    elif ext == ".pptx":
        doc = _pptx(data)
    elif ext in (".txt", ".md", ".markdown"):
        doc = ParsedDocument(data.decode("utf-8-sig", errors="replace"))
    elif ext in (".html", ".htm"):
        doc = _html(data)
    elif ext == ".rtf":
        doc = _rtf(data)
    elif ext == ".odt":
        doc = _odt(data)
    else:
        raise UnsupportedFile(f"Unsupported document type: {ext or filename}")
    doc.text = re.sub(r"\n{3,}", "\n\n", doc.text).strip()
    if len(doc.text) > MAX_TEXT_CHARS:
        doc.text = doc.text[:MAX_TEXT_CHARS]
        doc.warnings.append("Very long file: only the first ~150k tokens were read.")
    return doc


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_spreadsheet(filename: str, data: bytes) -> list[ParsedSheet]:
    ext = PurePosixPath(filename.lower()).suffix
    if ext == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
        return [ParsedSheet(PurePosixPath(filename).stem, [r for r in rows if any(r)])]
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = []
        for ws in wb.worksheets:
            rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(r)]
            if rows:
                width = max(i + 1 for r in rows for i, v in enumerate(r) if v)
                sheets.append(ParsedSheet(ws.title, [r[:width] + [""] * (width - len(r[:width])) for r in rows]))
        return sheets
    raise UnsupportedFile(f"Unsupported spreadsheet type: {ext or filename}")
