"""Reading notes and documents into the graph: the proposal, the review, and what is and isn't applied on its own.

The AI is a fake (conftest's fake_structured), so these check what the app does with an answer, not the answer."""
import io

import pytest
from crm_helpers import NOTES, entity, extraction, link, make, post_notes, upload

from core.server import ai, events, jobs, settings
from modules.crm.server import ingest, inbox
from modules.crm.server.extraction import (EXTRACTION_SYSTEM, ExtractedEntity, ExtractedRelationship, Extraction,
                                           Triage, extract_heuristic)

@pytest.fixture
def rules():
    """No AI at all: the plain-text reading."""
    settings.update({"crm.extraction_provider": "rules", "crm.utility_provider": "rules"})


# ── The rule-based reader ─────────────────────────────────────────────

def test_the_rule_based_reader_finds_attendees_organizations_and_action_items():
    ex = extract_heuristic(NOTES, owner_name="Me", known_entities=[("Frostline Logistics", "organization")])
    by_name = {e.name: e for e in ex.entities}
    assert ex.title == "Helio pilot sync"
    assert by_name["Priya Nair"].role == "attendee" and by_name["Tom Albright"].email == "tom@frostline.example"
    assert by_name["Me"].role == "attendee" and by_name["Frostline Logistics"].role == "mentioned"
    assert by_name["Helio Cold Chain"].kind == "organization"
    [rel] = ex.relationships
    assert rel.kind == "works_at"
    assert ex.action_items == ["Tom to send the fleet list", "Me to draft the SOW"]
    assert "pilot results" in ex.summary


def test_the_extraction_schema_is_flat_enough_for_claude():
    from core.server import ai_schema
    text = str(ai_schema.strict_schema(Extraction))
    assert "$ref" not in text and "anyOf" not in text


# ── From notes to the graph ───────────────────────────────────────────

def test_notes_become_a_proposal_you_review_and_then_apply(client, rules):
    me = client.put("/api/crm/me", json={"name": "Me"}).json()
    priya = make(client, "person", "Priya Nair")
    make(client, "organization", "Frostline Logistics")

    job = post_notes(client, occurred_at="2026-05-04T10:00:00Z")
    proposal = job["proposal"]
    assert job["status"] == "review" and job["extractor"] == "rules"
    assert proposal["doc_type"] == "meeting_notes" and proposal["record"]["kind"] == "meeting"
    ents = {e["name"]: e for e in proposal["entities"]}
    assert (ents["Priya Nair"]["action"], ents["Priya Nair"]["match_id"]) == ("link", priya["id"])
    assert ents["Me"]["match_id"] == me["id"] and ents["Me"]["is_self"]
    assert ents["Tom Albright"]["action"] == "create" and ents["Frostline Logistics"]["action"] == "link"

    for e in proposal["entities"]:   # the reviewer drops the organization mention
        if e["name"] == "Frostline Logistics":
            e["action"] = "skip"
    r = client.post(f"/api/crm/ingest/jobs/{job['id']}/apply", json={"proposal": proposal})
    assert r.status_code == 200, r.text
    result = r.json()["result"]
    assert (result["linked"], result["created"]) == (2, 3)   # Priya and Me; the meeting, Tom, Helio Cold Chain
    assert result["record_kind"] == "meeting"

    meeting = client.get(f"/api/crm/entities/{result['record_id']}").json()
    assert meeting["entity"]["occurred_at"].startswith("2026-05-04")
    assert meeting["entity"]["attributes"]["action_items"][0] == "Tom to send the fleet list"
    assert {r["other"]["name"] for r in meeting["relationships"] if r["kind"] == "attended"} == {"Priya Nair", "Tom Albright", "Me"}
    tom = client.get("/api/crm/entities", params={"q": "Tom Albright"}).json()["items"][0]
    assert tom["attributes"]["email"] == "tom@frostline.example"
    timeline = client.get(f"/api/crm/entities/{priya['id']}/timeline").json()
    assert timeline[0]["type"] == "meeting" and timeline[0]["title"] == "Helio pilot sync"
    notes = client.get(f"/api/crm/entities/{result['record_id']}/notes").json()
    assert any("Priya shared pilot results" in n["body"] for n in notes)   # the original text is kept

    assert client.post(f"/api/crm/ingest/jobs/{job['id']}/apply", json={"proposal": proposal}).status_code == 422


def test_an_apply_that_makes_no_sense_is_refused_and_the_item_stays_in_review(client, rules):
    job = post_notes(client, "Attendees: Zed Zulu\nQuick chat.")
    proposal = job["proposal"]
    proposal["entities"][0].update(action="link", match_id=None)
    assert client.post(f"/api/crm/ingest/jobs/{job['id']}/apply", json={"proposal": proposal}).status_code == 422
    assert client.get(f"/api/crm/ingest/jobs/{job['id']}").json()["status"] == "review"


def test_a_failed_apply_changes_nothing(client, rules, monkeypatch):
    job = post_notes(client, "Attendees: Zed Zulu, Amy Pond\nKickoff for the depot pilot.")
    monkeypatch.setattr("modules.crm.server.store.upsert_edge", lambda *a, **k: 1 / 0)
    with pytest.raises(ZeroDivisionError):
        ingest.apply_proposal(job["id"], job["proposal"])
    assert client.get("/api/crm/entities", params={"q": "Zed Zulu"}).json()["total"] == 0
    assert client.get("/api/crm/entities", params={"kind": "meeting"}).json()["total"] == 0
    assert inbox.require(job["id"])["status"] == "review"


def test_a_name_that_matches_several_people_is_never_linked_automatically(client, rules):
    make(client, "person", "Sarah Chen")
    make(client, "person", "Sarah Cohen")
    sarah = post_notes(client, "Attendees: Sarah\nIntro chat.")["proposal"]["entities"][0]
    assert sarah["action"] == "create" and sarah["ambiguous"] is True and len(sarah["candidates"]) == 2


def test_auto_apply_holds_ambiguous_names_and_does_not_make_duplicate_people(client, rules):
    make(client, "person", "Sarah Chen")
    make(client, "person", "Sarah Cohen")
    a = "Attendees: Zed Zulu, Sarah\nKickoff for the depot pilot. We agreed on scope, owners, and a date for the first site visit."
    b = "Attendees: Zed Zulu\nFollow-up call about the pilot. Zed confirmed budget and asked for the sensor spec sheet by Friday."
    done = upload(client, ("a.txt", a.encode()), ("b.txt", b.encode()))
    assert done["a.txt"]["status"] == "review"    # "Sarah" matched two people: held for a person
    assert done["b.txt"]["status"] == "applied"
    r = client.post(f"/api/crm/ingest/jobs/{done['a.txt']['id']}/apply", json={"proposal": done["a.txt"]["proposal"]})
    assert r.status_code == 200, r.text
    assert client.get("/api/crm/entities", params={"q": "Zed Zulu"}).json()["total"] == 1   # linked, not duplicated


def test_apply_checks_again_for_an_exact_match_when_it_creates(client, rules):
    job = post_notes(client, "Attendees: Zed Zulu\nQuick chat about the depot.")
    assert job["proposal"]["entities"][0]["action"] == "create"
    zed = make(client, "person", "Zed Zulu")   # appeared after the proposal was made
    client.post(f"/api/crm/ingest/jobs/{job['id']}/apply", json={"proposal": job["proposal"]})
    assert client.get("/api/crm/entities", params={"q": "Zed Zulu"}).json()["total"] == 1
    assert zed["id"] in inbox.require(job["id"])["result"]["entity_ids"].values()


# ── With an AI ────────────────────────────────────────────────────────

def test_the_ai_reading_goes_into_the_proposal_with_the_names_it_was_given(client, fake_structured):
    client.put("/api/crm/me", json={"name": "Me"})
    priya = make(client, "person", "Priya Nair")
    fake_structured.reply = extraction([
        entity("p1", "Priya Nair", role="attendee", title="CEO"),
        entity("o1", "Helio Cold Chain", kind="organization"),
        entity("t1", "cold chain logistics", kind="topic"),
    ], relationships=[ExtractedRelationship(source="p1", target="o1", kind="works_at", context="intro")],
        action_items=["Priya to send the report"], tags=["Pilot"])
    job = post_notes(client, "Spoke with Priya Nair about the Helio pilot and cold chain logistics. Plenty to do next.")
    assert job["status"] == "review" and job["extractor"].startswith("anthropic:")
    ents = {e["name"]: e for e in job["proposal"]["entities"]}
    assert ents["Priya Nair"]["match_id"] == priya["id"] and ents["Priya Nair"]["title"] == "CEO"
    assert job["proposal"]["record"]["tags"] == ["pilot"]
    assert {r["kind"] for r in job["proposal"]["relationships"]} == {"attended", "discussed_in", "works_at"}
    [call] = fake_structured.calls
    assert call["system"] == EXTRACTION_SYSTEM and call["schema"] is Extraction and call["choice"] == "anthropic"
    assert "Priya Nair" in call["prompt"].split("Known entities")[1]   # existing names are offered, to reuse the spelling


def test_a_meeting_gets_attended_for_people_and_discussed_in_for_everything_else(client, fake_structured):
    fake_structured.reply = extraction([entity("p1", "Ana", role="attendee"), entity("p2", "Raj", role="mentioned"),
                                       entity("o1", "Helio", kind="organization")])
    proposal = post_notes(client, "Ana and the team met to discuss Raj at Helio. A fairly long note.")["proposal"]
    kinds = {(r["source"], r["kind"]) for r in proposal["relationships"]}
    assert kinds == {("p1", "attended"), ("p2", "discussed_in"), ("o1", "discussed_in")}


def test_when_the_provider_is_down_the_batch_parks_and_resume_finishes_it_without_duplicates(client, fake_structured):
    down = ai.AIError("Claude (Anthropic): The Anthropic account is out of credit (Plans & Billing)", provider_down=True)
    fake_structured.reply = down
    files = [(f"n{i}.txt", f"Attendees: Zed Zulu, Person{i}\nNotes number {i} about the depot pilot, with enough words to count.".encode())
             for i in range(5)]
    first = upload(client, *files)
    assert all(j["status"] == "pending" for j in first.values())          # nothing was read, nothing is lost
    summary = client.get("/api/crm/ingest/summary").json()
    assert summary["parked"]["count"] == 5 and "out of credit" in summary["parked"]["reason"]
    assert "Resume" in summary["parked"]["reason"]
    assert len(fake_structured.calls) < 5          # once it was seen to be down, the rest weren't even tried
    assert all(j["work"]["status"] == "waiting" for j in first.values())
    assert client.get("/api/crm/entities", params={"kind": "person"}).json()["total"] == 0

    fake_structured.reply = lambda prompt, system, schema: extraction([entity("p1", "Zed Zulu", role="attendee")])
    assert client.post("/api/crm/ingest/resume").json()["resumed"] == 5
    jobs.drain()
    done = client.get("/api/crm/ingest/jobs", params={"status": "applied"}).json()
    assert len(done) == 5
    assert client.get("/api/crm/entities", params={"q": "Zed Zulu"}).json()["total"] == 1   # one Zed, not five


def test_changing_the_ai_settings_resumes_what_was_parked(client, fake_structured):
    fake_structured.reply = ai.AIError("Claude (Anthropic): Claude rejected the API key", provider_down=True)
    upload(client, ("a.txt", b"Attendees: Zed Zulu\nA note about the depot pilot with enough words in it to be read."))
    assert client.get("/api/crm/ingest/summary").json()["parked"]["count"] == 1
    fake_structured.reply = lambda prompt, system, schema: extraction([entity("p1", "Zed Zulu", role="attendee")])
    client.put("/api/settings", json={"ai.anthropic_api_key": "sk-new-key-12345678"})
    jobs.drain()
    assert client.get("/api/crm/ingest/summary").json()["counts"]["applied"] == 1


def test_a_reading_that_failed_on_this_text_falls_back_to_rules_and_waits_for_a_person(client, fake_structured):
    fake_structured.reply = ai.AIError("Claude declined to process this text")   # a content failure
    job = post_notes(client, "Attendees: Zed Zulu\nQuick chat about the depot pilot, with enough words to count here.",
                     title="Depot")
    assert job["status"] == "review" and job["extractor"] == "rules"
    assert job["proposal"]["warnings"][0].startswith(ingest.FALLBACK_WARNING_PREFIX)
    assert ingest.needs_review(job["proposal"])
    # Even with auto-apply on, a rules reading after an AI failure is held.
    held = upload(client, ("b.txt", b"Attendees: Amy Pond\nAnother chat about the depot pilot, with enough words to count here."))
    assert held["b.txt"]["status"] == "review"
    assert client.post("/api/crm/ingest/jobs/apply-ready").json() == {"applied": 0, "held": 2}


def test_apply_ready_applies_only_what_is_safe(client, rules):
    make(client, "person", "Sarah Chen")
    make(client, "person", "Sarah Cohen")
    upload(client, ("safe.txt", b"Attendees: Zed Zulu\nA clear note about the depot pilot, with enough words in it to count."),
           ("unsure.txt", b"Attendees: Sarah\nA vague note about the depot pilot, with enough words in it to count."),
           auto_apply=False)
    assert client.post("/api/crm/ingest/jobs/apply-ready").json() == {"applied": 1, "held": 1}


# ── Files ─────────────────────────────────────────────────────────────

def test_files_already_imported_or_discarded_or_duplicated_or_unreadable_are_skipped_with_a_reason(client, rules):
    text = b"Attendees: Zed Zulu\nA note about the depot pilot, with enough words in it to be worth reading."
    first = upload(client, ("a.txt", text), auto_apply=False)["a.txt"]
    again = upload(client, ("copy of a.txt", text))["copy of a.txt"]
    assert again["status"] == "skipped" and "Already imported" in again["error"]

    client.post(f"/api/crm/ingest/jobs/{first['id']}/discard")
    thrown = upload(client, ("a (1).txt", text))["a (1).txt"]
    assert thrown["status"] == "skipped" and "discarded" in thrown["error"]   # stays discarded

    other = b"Attendees: Amy Pond\nA different note about the depot pilot, with enough words in it to be worth reading."
    batch = upload(client, ("x.txt", other), ("y.txt", other), ("empty.txt", b"hi"), ("song.mp3", b"\x00"), ("data.xyz", b"?"))
    assert batch["y.txt"]["status"] == "skipped" and "Duplicate of x.txt" in batch["y.txt"]["error"]
    assert batch["empty.txt"]["error"] == "File has no usable text."
    assert "transcript" in batch["song.mp3"]["error"] and "Unsupported file type" in batch["data.xyz"]["error"]
    assert batch["x.txt"]["status"] == "applied"


def test_a_google_drive_zip_is_unpacked_and_other_google_formats_are_read(client, rules):
    import io
    import zipfile
    note = "Attendees: Zed Zulu\nA note about the depot pilot, with enough words in it to be worth reading, {}."
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:                       # a Drive folder download: Docs come as .docx, here as text
        z.writestr("Research/Zed Zulu.md", note.format("from the zip"))
        z.writestr("Research/__MACOSX/._Zed Zulu.md", "junk")
        z.writestr("Research/inner.zip", b"PK")
    html = f"<html><head><title>Depot</title><script>x()</script></head><body><p>{note.format('as a web page')}</p></body></html>"
    rtf = "{\\rtf1\\ansi{\\fonttbl{\\f0 Helvetica;}}" + note.format("in rich text").replace("\n", "\\par ") + "}"
    batch = upload(client, ("Drive/research.zip", buf.getvalue()), ("page.html", html.encode()), ("notes.rtf", rtf.encode()),
                   ("Zed.gdoc", b'{"doc_id": "abc"}'), ("broken.zip", b"not a zip"))
    assert batch["Zed Zulu.md"]["status"] == "applied" and batch["Zed Zulu.md"]["payload"]["folder"] == "Drive/Research"
    assert "__MACOSX" not in str(list(batch)) and "zip inside a zip" in batch["inner.zip"]["error"]
    assert batch["page.html"]["status"] == "applied" and "x()" not in client.get(f"/api/crm/ingest/jobs/{batch['page.html']['id']}").json()["raw_text"]
    assert batch["notes.rtf"]["status"] == "applied" and "fonttbl" not in batch["notes.rtf"]["raw_text"]
    assert batch["Zed.gdoc"]["status"] == "skipped" and "Google Drive shortcut" in batch["Zed.gdoc"]["error"]
    assert batch["broken.zip"]["status"] == "skipped" and "Couldn't open the zip" in batch["broken.zip"]["error"]


def test_folder_and_sibling_names_reach_the_ai(client, fake_structured):
    fake_structured.reply = lambda prompt, system, schema: extraction([entity("p1", "Zed Zulu", role="attendee")])
    upload(client, ("Clients/Acme/notes 9_11.txt", b"Attendees: Zed Zulu\nNotes from the Acme call, with enough words in them to count."),
           ("Clients/Acme/agenda.txt", b"Agenda for the Acme call with the team, with enough words in it to count too."))
    prompt = next(c["prompt"] for c in fake_structured.calls if "File name: notes 9_11.txt" in c["prompt"])  # they run side by side
    assert "File name: notes 9_11.txt" in prompt and "Folder: Clients/Acme" in prompt and "agenda.txt" in prompt


def test_a_long_file_is_sorted_first_and_a_saved_web_page_is_only_skimmed(client, fake_structured):
    settings.update({"crm.extraction_provider": "anthropic", "crm.utility_provider": "gemini"})
    verdicts = []

    def reply(prompt, system, schema):
        if schema is Triage:
            verdicts.append(prompt)
            return Triage(category="reference_material", process="brief", reason="A market report.")
        return extraction([], doc_type="document", title="Market report")

    fake_structured.reply = reply
    long_text = ("Cold chain market overview. " * 200).encode()
    job = upload(client, ("market www_site.txt", long_text))["market www_site.txt"]
    assert job["payload"]["triage"]["process"] == "brief" and len(verdicts) == 1
    assert job["status"] in ("review", "applied") and "Long reference" not in " ".join(job["proposal"]["warnings"])
    assert [c["choice"] for c in fake_structured.calls] == ["gemini", "anthropic"]   # sorting is the quick job

    junk = Triage(category="junk", process="skip", reason="Boilerplate.")
    fake_structured.reply = lambda prompt, system, schema: junk if schema is Triage else extraction([])
    skipped = upload(client, ("terms.txt", ("Terms and conditions. " * 150).encode()))["terms.txt"]
    assert skipped["status"] == "skipped" and skipped["error"] == "Boilerplate."


# ── Recorded meetings ─────────────────────────────────────────────────

GRANOLA = dict(origin="granola", external_id="gr-123", title="Pilot planning with Ana", started_at="2026-09-20T17:00:00Z",
               attendees=[{"name": "Me", "email": "me@example.com", "is_self": True},
                          {"name": "Ana Ruiz", "email": "ana@helio.example"}, {"email": "raj.patel@helio.example"}],
               transcript="Ana: thanks for making time. Raj: the pilot needs 40 fridges by March.\nMe: we can do that.")


def test_a_recorded_meeting_matches_invitees_by_email_and_is_not_imported_twice(client, rules):
    me = client.put("/api/crm/me", json={"name": "Me"}).json()
    ana = make(client, "person", "Ana R.", attributes={"email": "ana@helio.example"})
    row, seen = ingest.create_meeting_job(**GRANOLA)
    jobs.drain()
    row = inbox.require(row["id"])
    assert not seen and row["status"] == "applied", row
    ctx = client.get(f"/api/crm/entities/{row['result']['record_id']}/context").json()
    assert {(p["person"]["name"], p["relation"]) for p in ctx["people"]} == {("Me", "attended"), ("Ana R.", "attended"), ("Raj Patel", "attended")}
    assert {p["person"]["id"] for p in ctx["people"]} >= {ana["id"], me["id"]}
    detail = client.get(f"/api/crm/entities/{row['result']['record_id']}").json()["entity"]
    assert detail["attributes"]["external_ids"] == {"granola": "gr-123"}
    again, seen = ingest.create_meeting_job(**GRANOLA)
    assert seen and again["id"] == row["id"]
    resent, seen = ingest.create_meeting_job(**{**GRANOLA, "external_id": "gr-999"})   # same text, new id: still the same meeting
    assert seen and resent["id"] == row["id"]


def test_a_recorded_meeting_adds_to_the_same_meeting_you_took_notes_for(client, rules):
    client.put("/api/crm/me", json={"name": "Me"})
    ana = make(client, "person", "Ana Ruiz")
    earlier = make(client, "meeting", "Pilot planning", occurred_at="2026-09-20T16:00:00Z", summary="From my notes.")
    link(client, ana, "attended", earlier)
    row, _ = ingest.create_meeting_job(**{**GRANOLA, "external_id": "gr-456"})
    jobs.drain()
    row = inbox.require(row["id"])
    assert row["status"] == "applied" and row["result"]["record_id"] == earlier["id"]
    assert client.get("/api/crm/entities", params={"kind": "meeting"}).json()["total"] == 1
    kept = client.get(f"/api/crm/entities/{earlier['id']}").json()["entity"]
    assert kept["summary"] == "From my notes." and kept["attributes"]["external_ids"] == {"granola": "gr-456"}
    assert any("40 fridges" in n["body"] and n["source"] == "granola" for n in client.get(f"/api/crm/entities/{earlier['id']}/notes").json())


def test_a_recorded_meeting_needs_something_to_read():
    with pytest.raises(ingest.ProposalError):
        ingest.create_meeting_job(title="Empty")


def test_a_meetings_action_items_go_out_for_the_tasks_module_to_pick_up(client, fake_structured):
    got = []
    events.subscribe("actions.found", got.append)
    fake_structured.reply = extraction([entity("p1", "Zed Zulu", role="attendee")], action_items=["Zed to send the spec"])
    job = upload(client, ("a.txt", b"Attendees: Zed Zulu\nA note about the depot pilot, with enough words in it to count."))["a.txt"]
    assert job["status"] == "applied"
    [payload] = [p for p in got if p["source_ref"] == f"crm:{job['result']['record_id']}"]
    assert payload["items"] == ["Zed to send the spec"]
    assert any(t["text"] == "Zed to send the spec" for t in client.get("/api/tasks").json())

    # Enriching the same meeting later sends its whole list (the tasks module swaps a source's open tasks for it).
    got.clear()
    earlier = make(client, "meeting", "Depot chat", occurred_at="2026-10-01T10:00:00Z",
                   attributes={"action_items": ["Zed to send the spec"]})
    fake_structured.reply = extraction([entity("p1", "Zed Zulu", role="attendee")], title="Depot chat", occurred_on="2026-10-01",
                                       action_items=["Amy to book the van"])
    row = upload(client, ("b.txt", b"Attendees: Zed Zulu\nAnother depot note, with plenty more than enough words in it to count here as real text."))["b.txt"]
    assert row["result"]["record_id"] == earlier["id"] and [p["items"] for p in got] == [["Zed to send the spec", "Amy to book the van"]]


# ── Which lane, and who is asked ──────────────────────────────────────

def test_cloud_work_runs_four_at_a_time_and_local_or_rules_work_one_at_a_time(monkeypatch):
    monkeypatch.setattr(ai, "_ollama_tags", lambda cfg: [])
    assert ingest.lane("x") == "cloud"
    settings.update({"crm.extraction_provider": "rules"})
    assert ingest.lane("x") == "default"
    settings.update({"crm.extraction_provider": "ollama"})
    assert ingest.lane("x") == "default"


def test_a_reading_failure_that_is_not_the_providers_shows_on_the_item(client, rules, monkeypatch):
    monkeypatch.setattr("modules.crm.server.ingest.build_proposal", lambda *a, **k: 1 / 0)
    r = client.post("/api/crm/ingest/notes", json={"text": "Attendees: Zed Zulu\nQuick chat."})
    jobs.drain()
    row = client.get(f"/api/crm/ingest/jobs/{r.json()['id']}").json()
    assert row["status"] == "failed" and row["error"] == "division by zero" and row["work"]["status"] == "error"
    monkeypatch.undo()
    assert client.post(f"/api/crm/ingest/jobs/{row['id']}/retry").status_code == 202
    jobs.drain()
    assert client.get(f"/api/crm/ingest/jobs/{row['id']}").json()["status"] == "review"


def test_running_a_read_twice_does_nothing_the_second_time(client, rules):
    job = post_notes(client)
    before = inbox.require(job["id"])
    row_id = job["id"]
    ingest.handle_extract(jobs.Job("x", "crm", "extract", row_id, True, False))   # a restart replaying it
    assert inbox.require(row_id)["updated_at"] == before["updated_at"]


# ── Documents ─────────────────────────────────────────────────────────

def test_documents_are_read_to_text_including_tables_and_speaker_notes():
    import docx
    from pptx import Presentation

    from modules.crm.server.documents import UnsupportedFile, kind_of, parse_document, text_hash

    d = docx.Document()
    d.add_paragraph("Board update")
    t = d.add_table(rows=2, cols=2)
    for cell, text in zip(t._cells, ["Name", "Role", "Priya", "CEO"]):
        cell.text = text
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse_document("update.docx", buf.getvalue())
    assert "Board update" in parsed.text and "Priya | CEO" in parsed.text

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Pilot results"
    slide.notes_slide.notes_text_frame.text = "Mention the freezer trial"
    buf = io.BytesIO()
    p.save(buf)
    deck = parse_document("deck.pptx", buf.getvalue())
    assert "[Slide 1]" in deck.text and "Pilot results" in deck.text and "Speaker notes: Mention the freezer trial" in deck.text

    assert parse_document("n.md", b"# Hi\n\n\n\nThere").text == "# Hi\n\nThere"
    assert kind_of("a.XLSX") == "spreadsheet" and kind_of("a.mov") == "unsupported"
    with pytest.raises(UnsupportedFile):
        parse_document("a.zip", b"")
    assert text_hash("Same  text\n") == text_hash("same text")   # layout and case don't make a different document
